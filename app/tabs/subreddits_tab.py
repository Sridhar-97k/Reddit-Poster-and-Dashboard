import logging
from datetime import datetime

from PyQt5.QtWidgets import (QWidget, QVBoxLayout, QHBoxLayout, QGroupBox,
                             QLabel, QLineEdit, QPushButton, QTableWidget,
                             QTableWidgetItem, QHeaderView, QMessageBox,
                             QFileDialog, QApplication)
from PyQt5.QtGui import QFont
from PyQt5.QtCore import Qt
from openpyxl import Workbook
from openpyxl.styles import Font as XLFont, PatternFill, Alignment

logger = logging.getLogger(__name__)


class SubredditsTab(QWidget):
    def __init__(self, favorites_manager, get_reddit):
        super().__init__()
        self._favorites = favorites_manager
        self._get_reddit = get_reddit
        self._build_ui()
        self._refresh_table()

    def _build_ui(self):
        layout = QVBoxLayout(self)
        layout.setSpacing(20)

        info_group = QGroupBox("⭐ Favorite Subreddits")
        info_layout = QVBoxLayout()
        info_label = QLabel(
            "Manage your frequently-used subreddits here!\n\n"
            "Benefits:\n"
            "• Quick selection when creating Excel files\n"
            "• Auto-complete suggestions\n"
            "• One-click subreddit templates\n"
            "• Learn from your posting history"
        )
        info_label.setWordWrap(True)
        info_layout.addWidget(info_label)
        info_group.setLayout(info_layout)
        layout.addWidget(info_group)

        add_group = QGroupBox("➕ Add Favorite Subreddit")
        add_layout = QHBoxLayout()
        self.new_subreddit_input = QLineEdit()
        self.new_subreddit_input.setPlaceholderText("Enter subreddit name (without r/)")
        self.new_subreddit_input.returnPressed.connect(self._add_favorite)
        add_btn = QPushButton("➕ Add to Favorites")
        add_btn.clicked.connect(self._add_favorite)
        add_btn.setMinimumHeight(35)
        import_karma_btn = QPushButton("📥 Import from Karma Stats")
        import_karma_btn.clicked.connect(self._import_from_karma)
        import_karma_btn.setMinimumHeight(35)
        add_layout.addWidget(self.new_subreddit_input)
        add_layout.addWidget(add_btn)
        add_layout.addWidget(import_karma_btn)
        add_group.setLayout(add_layout)
        layout.addWidget(add_group)

        list_group = QGroupBox("📋 Your Favorite Subreddits")
        list_layout = QVBoxLayout()
        filter_layout = QHBoxLayout()
        filter_layout.addWidget(QLabel("Filter:"))
        self.filter_input = QLineEdit()
        self.filter_input.setPlaceholderText("Search favorites...")
        self.filter_input.textChanged.connect(self._refresh_table)
        filter_layout.addWidget(self.filter_input)
        list_layout.addLayout(filter_layout)

        self.favorites_table = QTableWidget()
        self.favorites_table.setColumnCount(4)
        self.favorites_table.setHorizontalHeaderLabels(["Subreddit", "Added", "Times Used", "Actions"])
        self.favorites_table.horizontalHeader().setSectionResizeMode(0, QHeaderView.Stretch)
        self.favorites_table.setAlternatingRowColors(True)
        self.favorites_table.setSelectionBehavior(QTableWidget.SelectRows)
        list_layout.addWidget(self.favorites_table)

        bulk_layout = QHBoxLayout()
        export_btn = QPushButton("📤 Export to File")
        export_btn.clicked.connect(self._export_favorites)
        import_file_btn = QPushButton("📥 Import from File")
        import_file_btn.clicked.connect(self._import_favorites_file)
        clear_btn = QPushButton("🗑️ Clear All")
        clear_btn.clicked.connect(self._clear_all)
        clear_btn.setProperty("class", "danger")
        bulk_layout.addWidget(export_btn)
        bulk_layout.addWidget(import_file_btn)
        bulk_layout.addStretch()
        bulk_layout.addWidget(clear_btn)
        list_layout.addLayout(bulk_layout)
        list_group.setLayout(list_layout)
        layout.addWidget(list_group)

        templates_group = QGroupBox("🚀 Quick Actions")
        templates_layout = QHBoxLayout()
        create_template_btn = QPushButton("📝 Create Excel Template with Favorites")
        create_template_btn.clicked.connect(self._create_template_with_favorites)
        create_template_btn.setMinimumHeight(40)
        copy_list_btn = QPushButton("📋 Copy List to Clipboard")
        copy_list_btn.clicked.connect(self._copy_to_clipboard)
        copy_list_btn.setMinimumHeight(40)
        templates_layout.addWidget(create_template_btn)
        templates_layout.addWidget(copy_list_btn)
        templates_group.setLayout(templates_layout)
        layout.addWidget(templates_group)

    def _refresh_table(self):
        filter_text = self.filter_input.text().lower() if hasattr(self, 'filter_input') else ''
        self.favorites_table.setRowCount(0)
        for fav in self._favorites.sorted_by_usage():
            name = fav['name']
            if filter_text and filter_text not in name:
                continue
            row = self.favorites_table.rowCount()
            self.favorites_table.insertRow(row)
            name_item = QTableWidgetItem(f"r/{name}")
            name_item.setFont(QFont('', 10, QFont.Bold))
            self.favorites_table.setItem(row, 0, name_item)
            added = datetime.fromisoformat(fav['added']).strftime("%Y-%m-%d")
            self.favorites_table.setItem(row, 1, QTableWidgetItem(added))
            usage_item = QTableWidgetItem(str(fav.get('times_used', 0)))
            usage_item.setTextAlignment(Qt.AlignCenter)
            self.favorites_table.setItem(row, 2, usage_item)
            del_btn = QPushButton("🗑️ Remove")
            del_btn.clicked.connect(lambda _, n=name: self._remove_favorite(n))
            self.favorites_table.setCellWidget(row, 3, del_btn)

    def _add_favorite(self):
        name = self.new_subreddit_input.text().strip()
        if not name:
            QMessageBox.warning(self, 'Input Required', 'Please enter a subreddit name')
            return
        if self._favorites.add(name):
            self._refresh_table()
            self.new_subreddit_input.clear()
            logger.info(f"Added r/{name} to favorites")
        else:
            QMessageBox.information(self, 'Already Exists', f'r/{name} is already in your favorites!')

    def _remove_favorite(self, name):
        self._favorites.remove(name)
        self._refresh_table()
        logger.info(f"Removed r/{name} from favorites")

    def _import_from_karma(self):
        reddit = self._get_reddit()
        if not reddit:
            return
        try:
            user = reddit.user.me()
            subs = {s.subreddit.display_name.lower() for s in user.submissions.new(limit=50)}
            added = sum(1 for s in subs if self._favorites.add(s))
            if added:
                self._refresh_table()
                QMessageBox.information(self, 'Import Complete',
                    f'Added {added} new subreddits!\nTotal: {len(self._favorites.favorites)}')
            else:
                QMessageBox.information(self, 'No New Subreddits',
                    'All subreddits from your recent posts are already in favorites!')
        except Exception as e:
            QMessageBox.critical(self, 'Error', f'Failed to import from karma:\n{str(e)}')

    def _export_favorites(self):
        if not self._favorites.favorites:
            QMessageBox.warning(self, 'No Favorites', 'You have no favorites to export!')
            return
        file_path, _ = QFileDialog.getSaveFileName(
            self, "Export Favorites",
            f"subreddit_favorites_{datetime.now().strftime('%Y%m%d')}.txt",
            "Text Files (*.txt)"
        )
        if file_path:
            try:
                with open(file_path, 'w') as f:
                    f.write("# My Favorite Subreddits\n")
                    f.write(f"# Exported: {datetime.now().strftime('%Y-%m-%d %H:%M')}\n\n")
                    for fav in sorted(self._favorites.favorites, key=lambda x: x['name']):
                        f.write(f"r/{fav['name']}\n")
                QMessageBox.information(self, 'Success', f'Favorites exported to:\n{file_path}')
            except Exception as e:
                QMessageBox.critical(self, 'Error', f'Failed to export:\n{str(e)}')

    def _import_favorites_file(self):
        file_path, _ = QFileDialog.getOpenFileName(
            self, "Import Favorites", "", "Text Files (*.txt);;All Files (*.*)"
        )
        if not file_path:
            return
        try:
            added = 0
            with open(file_path) as f:
                for line in f:
                    line = line.strip()
                    if not line or line.startswith('#'):
                        continue
                    if self._favorites.add(line):
                        added += 1
            if added:
                self._refresh_table()
                QMessageBox.information(self, 'Import Complete',
                    f'Added {added} new subreddits!\nTotal: {len(self._favorites.favorites)}')
            else:
                QMessageBox.information(self, 'No New Subreddits', 'All subreddits were already in favorites!')
        except Exception as e:
            QMessageBox.critical(self, 'Error', f'Failed to import:\n{str(e)}')

    def _clear_all(self):
        if not self._favorites.favorites:
            QMessageBox.information(self, 'No Favorites', 'You have no favorites to clear!')
            return
        reply = QMessageBox.question(
            self, 'Confirm Clear All',
            f'Remove all {len(self._favorites.favorites)} favorites? This cannot be undone!',
            QMessageBox.Yes | QMessageBox.No
        )
        if reply == QMessageBox.Yes:
            self._favorites.favorites.clear()
            self._favorites.save()
            self._refresh_table()

    def _copy_to_clipboard(self):
        if not self._favorites.favorites:
            QMessageBox.warning(self, 'No Favorites', 'You have no favorites to copy!')
            return
        text = '\n'.join(f"r/{f['name']}" for f in sorted(self._favorites.favorites, key=lambda x: x['name']))
        QApplication.clipboard().setText(text)
        QMessageBox.information(self, 'Copied!',
            f'Copied {len(self._favorites.favorites)} subreddits to clipboard!')

    def _create_template_with_favorites(self):
        if not self._favorites.favorites:
            QMessageBox.warning(self, 'No Favorites',
                'Add some favorite subreddits first!')
            return
        file_path, _ = QFileDialog.getSaveFileName(
            self, "Save Template with Favorites",
            f"reddit_posts_with_favorites_{datetime.now().strftime('%Y%m%d')}.xlsx",
            "Excel Files (*.xlsx)"
        )
        if not file_path:
            return
        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Reddit Posts"
            headers = ['subreddit', 'title', 'url', 'flair']
            for col, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col)
                cell.value = header
                cell.font = XLFont(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="3498DB", end_color="3498DB", fill_type="solid")
                cell.alignment = Alignment(horizontal='center')
            for idx, fav in enumerate(sorted(self._favorites.favorites, key=lambda x: x['name']), 2):
                ws.cell(row=idx, column=1, value=fav['name'])
                ws.cell(row=idx, column=2, value="[Your post title here]")
                ws.cell(row=idx, column=3, value="https://")
            ws.column_dimensions['A'].width = 20
            ws.column_dimensions['B'].width = 50
            ws.column_dimensions['C'].width = 60
            ws.column_dimensions['D'].width = 20
            wb.save(file_path)
            QMessageBox.information(self, 'Template Created!',
                f'Template created with {len(self._favorites.favorites)} favorites!\n\n{file_path}')
        except Exception as e:
            QMessageBox.critical(self, 'Error', f'Failed to create template:\n{str(e)}')
