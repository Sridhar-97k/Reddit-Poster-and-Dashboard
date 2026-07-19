import re
import logging
from datetime import datetime

from PyQt5.QtCore import QThread, pyqtSignal
from praw.exceptions import RedditAPIException
from openpyxl import load_workbook

logger = logging.getLogger(__name__)


class RedditWorker(QThread):
    """Submits a list of posts to Reddit one by one on a background thread."""
    progress = pyqtSignal(int, str)
    finished = pyqtSignal(dict)
    error = pyqtSignal(str)

    def __init__(self, reddit, posts_data, start_index=0, delay_secs=30):
        super().__init__()
        self.reddit = reddit
        self.posts_data = posts_data
        self.start_index = start_index
        self.delay_secs = delay_secs
        self.is_running = True

    def _interruptible_sleep(self, total_secs):
        """Sleep in 1-second chunks so stop() can interrupt a long wait."""
        for _ in range(total_secs):
            if not self.is_running:
                break
            self.msleep(1000)

    def _parse_ratelimit_wait(self, error_msg):
        """Return seconds to wait extracted from a RATELIMIT error message."""
        m = re.search(r'try again in (\d+) minute', error_msg, re.IGNORECASE)
        if m:
            return int(m.group(1)) * 60
        m = re.search(r'try again in (\d+) second', error_msg, re.IGNORECASE)
        if m:
            return int(m.group(1))
        return 600  # conservative fallback: 10 minutes

    def _try_submit(self, subreddit, title, url, flair_text, idx):
        """Submit one post and apply flair. Returns the submission object."""
        sub = self.reddit.subreddit(subreddit)
        submission = sub.submit(title=title, url=url)

        if flair_text:
            try:
                flair_choices = list(submission.flair.choices())
                matched = next(
                    (f for f in flair_choices
                     if f['flair_text'] and flair_text.lower() in f['flair_text'].lower()),
                    None
                )
                if matched:
                    submission.flair.select(matched['flair_template_id'])
                    self.progress.emit(idx + 1, f"Applied flair '{flair_text}' to r/{subreddit}")
                else:
                    self.progress.emit(idx + 1, f"⚠️ Flair '{flair_text}' not found in r/{subreddit}")
            except Exception as flair_error:
                self.progress.emit(idx + 1, f"⚠️ Flair error: {str(flair_error)[:50]}")

        return submission

    def run(self):
        results = {
            'successful': [],
            'failed': [],
            'total': len(self.posts_data)
        }

        for idx, post in enumerate(self.posts_data):
            if not self.is_running:
                break

            subreddit = post.get('subreddit', '').strip()
            title = post.get('title', '').strip()
            url = post.get('url', '').strip()
            flair_text = post.get('flair', '').strip()
            row_num = self.start_index + idx

            if not subreddit or not title or not url:
                results['failed'].append({
                    'row': row_num,
                    'reason': 'Missing required fields',
                    'post': post
                })
                self.progress.emit(idx + 1, f"Skipped row {row_num}: Missing data")
                continue

            try:
                submission = self._try_submit(subreddit, title, url, flair_text, idx)
                results['successful'].append({
                    'row': row_num,
                    'subreddit': subreddit,
                    'title': title,
                    'url': submission.url,
                    'permalink': f"https://reddit.com{submission.permalink}",
                    'id': submission.id
                })
                self.progress.emit(idx + 1, f"✅ Posted to r/{subreddit}: {title[:50]}...")

            except RedditAPIException as e:
                is_ratelimit = any(
                    getattr(item, 'error_type', '') == 'RATELIMIT' for item in e.items
                )
                if is_ratelimit:
                    wait_secs = self._parse_ratelimit_wait(str(e))
                    self.progress.emit(idx + 1, f"⏳ Rate limited — waiting {wait_secs}s then retrying r/{subreddit}...")
                    self._interruptible_sleep(wait_secs)
                    if not self.is_running:
                        break
                    try:
                        submission = self._try_submit(subreddit, title, url, flair_text, idx)
                        results['successful'].append({
                            'row': row_num,
                            'subreddit': subreddit,
                            'title': title,
                            'url': submission.url,
                            'permalink': f"https://reddit.com{submission.permalink}",
                            'id': submission.id
                        })
                        self.progress.emit(idx + 1, f"✅ Posted to r/{subreddit} (after rate limit wait): {title[:50]}...")
                    except Exception as retry_e:
                        results['failed'].append({'row': row_num, 'reason': str(retry_e), 'post': post})
                        self.progress.emit(idx + 1, f"❌ Retry failed row {row_num}: {str(retry_e)[:50]}")
                else:
                    results['failed'].append({'row': row_num, 'reason': str(e), 'post': post})
                    self.progress.emit(idx + 1, f"❌ Failed row {row_num}: {str(e)[:50]}")

            except Exception as e:
                results['failed'].append({'row': row_num, 'reason': str(e), 'post': post})
                self.progress.emit(idx + 1, f"❌ Error row {row_num}: {str(e)[:50]}")

            if self.is_running and idx < len(self.posts_data) - 1:
                self.progress.emit(idx + 1, f"⏱️ Waiting {self.delay_secs}s before next post...")
                self._interruptible_sleep(self.delay_secs)

        self.finished.emit(results)

    def stop(self):
        self.is_running = False


class KarmaWorker(QThread):
    """Fetches the user's recent submissions on a background thread."""
    finished = pyqtSignal(list)
    error = pyqtSignal(str)

    def __init__(self, reddit, num_posts):
        super().__init__()
        self.reddit = reddit
        self.num_posts = num_posts

    def run(self):
        try:
            logger.info(f"KarmaWorker: fetching {self.num_posts} posts")
            user = self.reddit.user.me()
            posts = []
            for idx, submission in enumerate(user.submissions.new(limit=self.num_posts), 1):
                logger.debug(f"KarmaWorker: fetched {idx}/{self.num_posts}")
                posts.append({
                    'subreddit': submission.subreddit.display_name,
                    'title': submission.title,
                    'score': submission.score,
                    'comments': submission.num_comments,
                    'created': datetime.fromtimestamp(submission.created_utc),
                    'url': f"https://reddit.com{submission.permalink}",
                    'id': submission.id
                })
            logger.info(f"KarmaWorker: done — {len(posts)} posts")
            self.finished.emit(posts)
        except Exception as e:
            logger.error(f"KarmaWorker: {e}", exc_info=True)
            self.error.emit(str(e))


class BulkKarmaUpdateWorker(QThread):
    """Updates karma for all posts in the log file on a background thread."""
    progress = pyqtSignal(int, int)   # (updated_count, failed_count)
    finished = pyqtSignal(int, int)
    error = pyqtSignal(str)

    def __init__(self, reddit, post_log_file):
        super().__init__()
        self.reddit = reddit
        self.post_log_file = post_log_file

    def run(self):
        try:
            wb = load_workbook(self.post_log_file)
            ws = wb.active
            updated_count = 0
            failed_count = 0

            for row in range(2, ws.max_row + 1):
                post_id = ws.cell(row=row, column=1).value
                if not post_id:
                    continue
                try:
                    submission = self.reddit.submission(id=post_id)
                    ws.cell(row=row, column=8, value=submission.score)
                    ws.cell(row=row, column=9, value=submission.num_comments)
                    ws.cell(row=row, column=10, value=datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
                    updated_count += 1
                    logger.debug(f"BulkKarmaUpdateWorker: updated {post_id}")
                except Exception as e:
                    logger.warning(f"BulkKarmaUpdateWorker: failed {post_id}: {e}")
                    failed_count += 1
                self.progress.emit(updated_count, failed_count)

            wb.save(self.post_log_file)
            logger.info(f"BulkKarmaUpdateWorker: done — {updated_count} updated, {failed_count} failed")
            self.finished.emit(updated_count, failed_count)
        except Exception as e:
            logger.error(f"BulkKarmaUpdateWorker: {e}", exc_info=True)
            self.error.emit(str(e))
