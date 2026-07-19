import os
import sys
import logging
from datetime import datetime

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment

logger = logging.getLogger(__name__)

HEADERS = [
    'Post ID', 'Posted Date', 'Posted Time', 'Subreddit',
    'Title', 'URL', 'Reddit Link', 'Score', 'Comments',
    'Last Updated', 'Status', 'Notes'
]
COL_WIDTHS = [15, 12, 10, 18, 50, 40, 50, 10, 10, 18, 12, 30]


class PostLog:
    def __init__(self, filepath='reddit_posts_log.xlsx'):
        self.filepath = filepath

    def initialize(self):
        """Create the log file with headers if it doesn't already exist."""
        if os.path.exists(self.filepath):
            logger.info(f"Post log exists: {self.filepath}")
            return

        try:
            wb = Workbook()
            ws = wb.active
            ws.title = "Post Log"

            for col, (header, width) in enumerate(zip(HEADERS, COL_WIDTHS), 1):
                cell = ws.cell(row=1, column=col)
                cell.value = header
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
                cell.alignment = Alignment(horizontal='center')
                ws.column_dimensions[cell.column_letter].width = width

            wb.save(self.filepath)
            logger.info(f"Created post log: {self.filepath}")
        except Exception as e:
            logger.error(f"Failed to create post log: {e}")

    def log_post(self, post_data):
        """Append one successfully submitted post to the log."""
        try:
            wb = load_workbook(self.filepath)
            ws = wb.active
            row = ws.max_row + 1
            now = datetime.now()

            ws.cell(row=row, column=1, value=post_data.get('id', ''))
            ws.cell(row=row, column=2, value=now.strftime('%Y-%m-%d'))
            ws.cell(row=row, column=3, value=now.strftime('%H:%M:%S'))
            ws.cell(row=row, column=4, value=post_data.get('subreddit', ''))
            ws.cell(row=row, column=5, value=post_data.get('title', ''))
            ws.cell(row=row, column=6, value=post_data.get('url', ''))
            ws.cell(row=row, column=7, value=post_data.get('permalink', ''))
            ws.cell(row=row, column=8, value=post_data.get('score', 0))
            ws.cell(row=row, column=9, value=post_data.get('comments', 0))
            ws.cell(row=row, column=10, value=now.strftime('%Y-%m-%d %H:%M:%S'))
            ws.cell(row=row, column=11, value='Posted')
            ws.cell(row=row, column=12, value='')

            for col in (8, 9):
                ws.cell(row=row, column=col).alignment = Alignment(horizontal='center')

            wb.save(self.filepath)
            logger.info(f"Logged post at row {row}: {post_data.get('title', '')[:40]}")
            return True
        except Exception as e:
            logger.error(f"Failed to log post: {e}")
            return False

    def update_karma(self, post_id, score, comments):
        """Update score and comment count for a specific post ID."""
        try:
            wb = load_workbook(self.filepath)
            ws = wb.active
            for row in range(2, ws.max_row + 1):
                if ws.cell(row=row, column=1).value == post_id:
                    ws.cell(row=row, column=8, value=score)
                    ws.cell(row=row, column=9, value=comments)
                    ws.cell(row=row, column=10, value=datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
                    wb.save(self.filepath)
                    return True
            logger.warning(f"Post {post_id} not found in log")
            return False
        except Exception as e:
            logger.error(f"Failed to update karma: {e}")
            return False

    def get_stats(self):
        """Return aggregate stats from the log, or None on error."""
        if not os.path.exists(self.filepath):
            return None
        try:
            wb = load_workbook(self.filepath, data_only=True)
            ws = wb.active
            total_posts = ws.max_row - 1
            total_karma = sum((ws.cell(row=r, column=8).value or 0) for r in range(2, ws.max_row + 1))
            total_comments = sum((ws.cell(row=r, column=9).value or 0) for r in range(2, ws.max_row + 1))
            return {
                'total_posts': total_posts,
                'total_karma': total_karma,
                'total_comments': total_comments,
                'avg_karma': total_karma / total_posts if total_posts > 0 else 0,
            }
        except Exception as e:
            logger.error(f"Failed to get log stats: {e}")
            return None

    def open_file(self):
        """Open the log file in the system default application."""
        if not os.path.exists(self.filepath):
            return False
        try:
            if sys.platform == 'win32':
                os.startfile(self.filepath)
            elif sys.platform == 'darwin':
                os.system(f'open "{self.filepath}"')
            else:
                os.system(f'xdg-open "{self.filepath}"')
            return True
        except Exception as e:
            logger.error(f"Failed to open log file: {e}")
            return False
