import logging
from datetime import datetime

from PyQt5.QtWidgets import QFileDialog, QMessageBox
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

logger = logging.getLogger(__name__)


def export_table_to_excel(parent, table, sheet_title, dialog_title, filename_prefix, col_widths=None):
    """Save a QTableWidget's contents to a user-chosen .xlsx file."""
    file_path, _ = QFileDialog.getSaveFileName(
        parent,
        dialog_title,
        f"{filename_prefix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx",
        "Excel Files (*.xlsx)"
    )
    if not file_path:
        return

    try:
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_title

        col_count = table.columnCount()
        for col in range(col_count):
            header_item = table.horizontalHeaderItem(col)
            cell = ws.cell(row=1, column=col + 1)
            cell.value = header_item.text() if header_item else f"Column {col + 1}"
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="3498DB", end_color="3498DB", fill_type="solid")

        for row in range(table.rowCount()):
            for col in range(col_count):
                item = table.item(row, col)
                ws.cell(row=row + 2, column=col + 1, value=item.text() if item else '')

        if col_widths:
            for col_letter, width in col_widths.items():
                ws.column_dimensions[col_letter].width = width

        wb.save(file_path)
        QMessageBox.information(parent, 'Success', f'Exported successfully!\n\n{file_path}')
        logger.info(f"Exported {sheet_title} to {file_path}")

    except Exception as e:
        QMessageBox.critical(parent, 'Error', f'Failed to export:\n{str(e)}')
